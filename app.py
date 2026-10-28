"""Import every CSV and Excel dataset in data/ into MySQL."""

from __future__ import annotations

import argparse
import csv
import codecs
import logging
import math
import os
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path
from typing import Iterator, Sequence

import mysql.connector
from dotenv import load_dotenv
from mysql.connector import Error
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_DATA_DIR = PROJECT_DIR / "data"
BATCH_SIZE = 1_000
CSV_ENCODINGS = ("utf-8-sig", "cp1252")
LOGGER = logging.getLogger("mysql_import")


@dataclass(frozen=True)
class Dataset:
    path: Path
    table_name: str
    sheet_name: str | None = None


@lru_cache(maxsize=None)
def csv_encoding(path: Path) -> str:
    """Choose UTF-8 or Windows-1252 without loading the CSV into memory."""
    for encoding in CSV_ENCODINGS:
        decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
        try:
            with path.open("rb") as file:
                for chunk in iter(lambda: file.read(64 * 1024), b""):
                    decoder.decode(chunk)
                decoder.decode(b"", final=True)
            return encoding
        except UnicodeDecodeError:
            continue
    raise ValueError(
        f"{path.name} is not valid UTF-8 or Windows-1252; convert it to UTF-8 "
        "before importing."
    )


def safe_identifier(value: str) -> str:
    """Turn a source name into a predictable SQL identifier."""
    identifier = re.sub(r"[^a-zA-Z0-9_]+", "_", value.strip()).strip("_").lower()
    if not identifier:
        raise ValueError(f"Cannot create a SQL name from {value!r}.")
    if identifier[0].isdigit():
        identifier = f"_{identifier}"
    return identifier


def unique_column_names(headers: Sequence[object]) -> list[str]:
    names: list[str] = []
    used: set[str] = set()
    for index, header in enumerate(headers, start=1):
        name = (
            safe_identifier(str(header))
            if header is not None and str(header).strip()
            else f"column_{index}"
        )
        candidate = name
        suffix = 2
        while candidate in used:
            candidate = f"{name}_{suffix}"
            suffix += 1
        names.append(candidate)
        used.add(candidate)
    if not names:
        raise ValueError("A dataset has no column headers.")
    return names


def discover_datasets(data_dir: Path) -> list[Dataset]:
    if not data_dir.is_dir():
        raise FileNotFoundError(f"Data folder not found: {data_dir}")

    datasets: list[Dataset] = []
    for path in sorted(data_dir.iterdir()):
        if not path.is_file():
            continue
        extension = path.suffix.lower()
        if extension == ".csv":
            datasets.append(Dataset(path=path, table_name=safe_identifier(path.stem)))
        elif extension == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet_names = workbook.sheetnames
            finally:
                workbook.close()
            if not sheet_names:
                raise ValueError(f"Excel workbook has no worksheets: {path.name}")
            for sheet_name in sheet_names:
                suffix = f"_{sheet_name}" if len(sheet_names) > 1 else ""
                datasets.append(
                    Dataset(
                        path=path,
                        table_name=safe_identifier(f"{path.stem}{suffix}"),
                        sheet_name=sheet_name,
                    )
                )

    if not datasets:
        raise FileNotFoundError(f"No .csv or .xlsx files found in {data_dir}")

    table_names = [dataset.table_name for dataset in datasets]
    if len(table_names) != len(set(table_names)):
        raise ValueError("Multiple input files resolve to the same MySQL table name.")
    return datasets


def get_headers(dataset: Dataset) -> list[str]:
    if dataset.path.suffix.lower() == ".csv":
        with dataset.path.open("r", encoding=csv_encoding(dataset.path), newline="") as file:
            headers = next(csv.reader(file), None)
    else:
        workbook = load_workbook(dataset.path, read_only=True, data_only=True)
        try:
            worksheet = workbook[dataset.sheet_name]
            headers = next(worksheet.iter_rows(values_only=True), None)
        finally:
            workbook.close()

    if headers is None:
        raise ValueError(f"Dataset is empty: {dataset.path.name}")
    return unique_column_names(headers)


def iter_dataset_rows(dataset: Dataset) -> Iterator[Sequence[object]]:
    if dataset.path.suffix.lower() == ".csv":
        with dataset.path.open("r", encoding=csv_encoding(dataset.path), newline="") as file:
            reader = csv.reader(file)
            next(reader, None)
            for row in reader:
                if row and any(value.strip() for value in row):
                    yield row
        return

    workbook = load_workbook(dataset.path, read_only=True, data_only=True)
    try:
        worksheet = workbook[dataset.sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if row and any(value is not None and str(value).strip() for value in row):
                yield row
    finally:
        workbook.close()


def numeric_kind(value: object) -> str:
    if isinstance(value, bool):
        return "text"
    if isinstance(value, (int, float, Decimal)):
        return "integer" if isinstance(value, int) or (
            isinstance(value, float) and value.is_integer()
        ) else "decimal"

    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return "text"
    if not number.is_finite():
        return "text"
    return "integer" if number == number.to_integral_value() else "decimal"


def infer_column_types(dataset: Dataset, column_count: int) -> list[str]:
    types = ["integer"] * column_count
    for row_number, row in enumerate(iter_dataset_rows(dataset), start=2):
        if len(row) > column_count:
            raise ValueError(
                f"{dataset.path.name}, row {row_number}: found {len(row)} values "
                f"but the header has only {column_count} columns."
            )
        for index, value in enumerate(row):
            if value is None or not str(value).strip():
                continue
            current = types[index]
            if current == "text":
                continue
            kind = numeric_kind(value)
            if kind == "text":
                types[index] = "text"
            elif kind == "decimal" and current == "integer":
                types[index] = "decimal"
    return types


def mysql_value(value: object, column_type: str) -> object:
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    if column_type == "integer":
        return int(value)
    if column_type == "decimal":
        number = float(value)
        if not math.isfinite(number):
            raise ValueError(f"Non-finite numeric value cannot be imported: {value!r}")
        return number
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def mysql_column_type(column_type: str) -> str:
    return {"integer": "BIGINT", "decimal": "DOUBLE", "text": "LONGTEXT"}[column_type]


def table_columns(connection: mysql.connector.MySQLConnection, table_name: str) -> list[str]:
    cursor = connection.cursor()
    try:
        cursor.execute(f"SHOW COLUMNS FROM `{table_name}`")
        return [row[0] for row in cursor.fetchall()]
    finally:
        cursor.close()


def import_dataset(connection: mysql.connector.MySQLConnection, dataset: Dataset) -> int:
    columns = get_headers(dataset)
    column_types = infer_column_types(dataset, len(columns))
    table = f"`{dataset.table_name}`"
    schema = ", ".join(
        f"`{column}` {mysql_column_type(column_type)} NULL"
        for column, column_type in zip(columns, column_types)
    )

    cursor = connection.cursor()
    try:
        cursor.execute(
            f"CREATE TABLE IF NOT EXISTS {table} "
            f"({schema}) ENGINE=InnoDB DEFAULT CHARACTER SET=utf8mb4"
        )
        existing_columns = table_columns(connection, dataset.table_name)
        if existing_columns != columns:
            raise ValueError(
                f"MySQL table {dataset.table_name!r} already exists with different "
                "columns. Rename or remove that table before importing."
            )

        placeholders = ", ".join(["%s"] * len(columns))
        column_list = ", ".join(f"`{column}`" for column in columns)
        insert_sql = f"INSERT INTO {table} ({column_list}) VALUES ({placeholders})"
        cursor.execute(f"DELETE FROM {table}")

        batch: list[tuple[object, ...]] = []
        imported_rows = 0
        for row_number, row in enumerate(iter_dataset_rows(dataset), start=2):
            if len(row) > len(columns):
                raise ValueError(
                    f"{dataset.path.name}, row {row_number}: found {len(row)} values "
                    f"but the header has only {len(columns)} columns."
                )
            padded_row = list(row) + [None] * (len(columns) - len(row))
            batch.append(
                tuple(
                    mysql_value(value, column_type)
                    for value, column_type in zip(padded_row, column_types)
                )
            )
            if len(batch) >= BATCH_SIZE:
                cursor.executemany(insert_sql, batch)
                imported_rows += len(batch)
                batch.clear()

        if batch:
            cursor.executemany(insert_sql, batch)
            imported_rows += len(batch)
        connection.commit()
        return imported_rows
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()


def required_setting(name: str, default: str | None = None) -> str:
    value = os.getenv(name, default)
    if value is None or not value.strip():
        raise ValueError(f"Set {name} in the project .env file before running the import.")
    return value.strip()


def run_import(data_dir: Path) -> None:
    database = required_setting("MYSQL_DATABASE", "air_quality_analysis")
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", database):
        raise ValueError("MYSQL_DATABASE must contain only letters, numbers, and underscores.")

    host = required_setting("MYSQL_HOST", "127.0.0.1")
    user = required_setting("MYSQL_USER")
    password = os.getenv("MYSQL_PASSWORD", "")
    try:
        port = int(required_setting("MYSQL_PORT", "3306"))
    except ValueError as error:
        raise ValueError("MYSQL_PORT must be a valid port number.") from error
    if not 1 <= port <= 65535:
        raise ValueError("MYSQL_PORT must be between 1 and 65535.")

    datasets = discover_datasets(data_dir)
    connection_options = {
        "host": host,
        "port": port,
        "user": user,
        "password": password,
        "connection_timeout": 15,
    }

    server = mysql.connector.connect(**connection_options)
    try:
        cursor = server.cursor()
        try:
            cursor.execute(
                f"CREATE DATABASE IF NOT EXISTS `{database}` "
                "CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci"
            )
            server.commit()
        finally:
            cursor.close()
    finally:
        server.close()

    connection = mysql.connector.connect(database=database, **connection_options)
    try:
        for dataset in datasets:
            count = import_dataset(connection, dataset)
            LOGGER.info(
                "Imported %s rows from %s%s into %s.%s",
                count,
                dataset.path.name,
                f" ({dataset.sheet_name})" if dataset.sheet_name else "",
                database,
                dataset.table_name,
            )
    finally:
        connection.close()


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    load_dotenv(PROJECT_DIR / ".env")

    parser = argparse.ArgumentParser(
        description="Import every .csv and .xlsx file in the data folder into MySQL."
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=DEFAULT_DATA_DIR,
        help="Folder containing CSV and Excel files (default: ./data).",
    )
    args = parser.parse_args()

    try:
        run_import(args.data_dir.resolve())
    except (Error, OSError, ValueError) as error:
        LOGGER.error("Import failed: %s", error)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())